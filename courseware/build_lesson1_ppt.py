#!/usr/bin/env python3
"""生成《纸上沉浮》第一课可编辑 PPT（16:9，面向初中课堂）。"""

from __future__ import annotations

import math
from pathlib import Path

import numpy as np
from lxml import etree
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
FONTS = ROOT / "fonts"
OUT_PPTX = ROOT / "纸上沉浮-第一课-初识诗仙李白.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
W_IN, H_IN = 13.333, 7.5

# 课堂电脑常见中文字体：标题用楷体，正文用微软雅黑
FONT_TITLE = "楷体"
FONT_BODY = "微软雅黑"
FONT_POETRY = "楷体"

INK = RGBColor(0x2B, 0x27, 0x23)
INK_SOFT = RGBColor(0x5C, 0x54, 0x4A)
CINNABAR = RGBColor(0xB1, 0x3A, 0x2C)
OCHRE = RGBColor(0xB0, 0x88, 0x4A)
INDIGO = RGBColor(0x3E, 0x52, 0x63)
MOSS = RGBColor(0x5A, 0x6B, 0x4E)
CARD = RGBColor(0xFF, 0xF9, 0xEF)
CARD_WARM = RGBColor(0xF7, 0xEB, 0xD5)
BORDER = RGBColor(0xD0, 0xBE, 0x9C)
WHITE = RGBColor(0xFF, 0xFC, 0xF6)
PAPER_LINE = RGBColor(0xC8, 0xB4, 0x8E)

KAI_TTF = FONTS / "LXGWWenKai-Medium.ttf"
KAI_REG = FONTS / "LXGWWenKai-Regular.ttf"
HEI_TTF = Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc")


def _font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    if path.exists():
        return ImageFont.truetype(str(path), size=size)
    if HEI_TTF.exists():
        return ImageFont.truetype(str(HEI_TTF), size=size)
    return ImageFont.load_default()


def make_paper(w: int = 1920, h: int = 1080, seed: int = 7) -> Image.Image:
    rng = np.random.default_rng(seed)
    base = np.array([246.0, 237.0, 218.0])
    arr = np.ones((h, w, 3), dtype=np.float32) * base
    arr += rng.normal(0, 5.8, (h, w, 1))
    # 横向纤维
    for _ in range(55):
        y0 = int(rng.integers(0, h - 2))
        x0 = int(rng.integers(0, w))
        length = int(rng.integers(90, 520))
        val = float(rng.normal(-7, 3.5))
        arr[y0 : y0 + 2, x0 : min(w, x0 + length)] += val
    # 淡墨渍
    yy, xx = np.mgrid[0:h, 0:w]
    for _ in range(7):
        cx, cy = int(rng.integers(0, w)), int(rng.integers(0, h))
        rx = float(rng.integers(80, 280))
        ry = float(rng.integers(50, 180))
        strength = float(rng.uniform(4, 12))
        mask = np.exp(-(((xx - cx) / rx) ** 2 + ((yy - cy) / ry) ** 2))
        arr -= mask[:, :, None] * strength
    arr = np.clip(arr, 0, 255).astype(np.uint8)
    img = Image.fromarray(arr, "RGB")
    return img


def _ridge(w: int, h: int, base_y: float, amp: float, seed: int) -> list[tuple[int, int]]:
    rng = np.random.default_rng(seed)
    x = np.arange(w, dtype=np.float64)
    y = np.zeros(w, dtype=np.float64)
    freq, amplitude = 0.0024, amp
    phase0 = rng.random() * math.tau
    for _ in range(5):
        phase = phase0 + rng.random() * 4
        y += np.sin(x * freq + phase) * amplitude
        y += np.sin(x * freq * 2.35 + phase * 1.4) * amplitude * 0.32
        freq *= 1.85
        amplitude *= 0.48
    y = (y - y.min()) / (y.max() - y.min() + 1e-9)
    y = base_y - y * amp
    pts = [(0, h)] + [(int(i), int(v)) for i, v in enumerate(y)] + [(w - 1, h)]
    return pts


def paint_mountains(img: Image.Image, strong: bool = False) -> Image.Image:
    w, h = img.size
    overlay = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    layers = (
        [
            (h * 0.62, 140, (112, 118, 122, 38), 14, 11),
            (h * 0.70, 170, (86, 92, 98, 52), 9, 22),
            (h * 0.78, 190, (62, 66, 72, 70), 6, 33),
            (h * 0.86, 150, (46, 48, 52, 92), 3, 44),
            (h * 0.93, 90, (36, 36, 38, 110), 2, 55),
        ]
        if strong
        else [
            (h * 0.78, 90, (120, 124, 126, 28), 12, 8),
            (h * 0.84, 110, (88, 92, 96, 40), 7, 18),
            (h * 0.90, 80, (58, 60, 64, 58), 4, 28),
            (h * 0.95, 50, (42, 42, 44, 72), 2, 38),
        ]
    )
    for base_y, amp, color, blur, seed in layers:
        layer = Image.new("RGBA", (w, h), (0, 0, 0, 0))
        draw = ImageDraw.Draw(layer)
        pts = _ridge(w, h, base_y, amp, seed)
        draw.polygon(pts, fill=color)
        if blur:
            layer = layer.filter(ImageFilter.GaussianBlur(blur))
        overlay = Image.alpha_composite(overlay, layer)
    # 远山淡雾
    fog = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    fd = ImageDraw.Draw(fog)
    fd.rectangle([0, int(h * 0.55), w, h], fill=(246, 237, 218, 35 if strong else 22))
    fog = fog.filter(ImageFilter.GaussianBlur(18))
    overlay = Image.alpha_composite(overlay, fog)
    return Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")


def make_brush_stroke(path: Path) -> None:
    w, h = 1400, 36
    img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    rng = np.random.default_rng(3)
    y = h / 2
    for x in range(w):
        t = x / w
        thick = 3.2 + 4.2 * math.sin(t * math.pi) ** 1.35
        thick *= 0.8 + 0.2 * rng.random()
        alpha = int(190 * math.sin(t * math.pi) ** 0.55)
        draw.ellipse(
            [x - 1, y - thick, x + 1, y + thick],
            fill=(43, 39, 35, max(0, min(255, alpha))),
        )
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    img.save(path)


def make_seal(text: str, path: Path, size: int = 280) -> None:
    img = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    m = 10
    red = (177, 58, 44, 235)
    draw.rounded_rectangle([m, m, size - m, size - m], radius=8, fill=red)
    draw.rounded_rectangle(
        [m + 9, m + 9, size - m - 9, size - m - 9],
        radius=4,
        outline=(255, 248, 236, 230),
        width=3,
    )
    chars = list(text)
    font = _font(KAI_TTF if KAI_TTF.exists() else KAI_REG, 78 if len(chars) <= 4 else 62)
    if len(chars) == 4:
        # 传统印章：右起直排。纸上沉浮 → 右列「纸/上」，左列「沉/浮」
        positions = [(1, 0), (1, 1), (0, 0), (0, 1)]
        cell = (size - 2 * m - 28) / 2
        origin = m + 18
        for ch, (gx, gy) in zip(chars, positions):
            cx = origin + gx * cell + cell / 2
            cy = origin + gy * cell + cell / 2
            bbox = draw.textbbox((0, 0), ch, font=font)
            tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
            draw.text((cx - tw / 2 - bbox[0], cy - th / 2 - bbox[1] - 4), ch, font=font, fill=(255, 248, 236, 255))
    else:
        font = _font(KAI_TTF, 92)
        bbox = draw.textbbox((0, 0), text, font=font)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.text(((size - tw) / 2 - bbox[0], (size - th) / 2 - bbox[1]), text, font=font, fill=(255, 248, 236, 255))
    img.save(path)


def build_assets() -> dict[str, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    cover = paint_mountains(make_paper(1920, 1080, 7), strong=True)
    inner = paint_mountains(make_paper(1920, 1080, 21), strong=False)
    close = paint_mountains(make_paper(1920, 1080, 13), strong=True)
    p_cover, p_inner, p_close = ASSETS / "bg_cover.png", ASSETS / "bg_inner.png", ASSETS / "bg_close.png"
    cover.save(p_cover, quality=92)
    inner.save(p_inner, quality=92)
    close.save(p_close, quality=92)
    stroke = ASSETS / "brush.png"
    seal = ASSETS / "seal.png"
    make_brush_stroke(stroke)
    make_seal("纸上沉浮", seal)
    return {"cover": p_cover, "inner": p_inner, "close": p_close, "stroke": stroke, "seal": seal}


def set_run_font(run, name: str, size_pt: float, color: RGBColor, bold: bool = False, spacing: int | None = None):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    if spacing is not None:
        rPr.set("spc", str(spacing))
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def add_text(
    slide,
    l,
    t,
    w,
    h,
    lines: list[str] | str,
    *,
    font=FONT_BODY,
    size=18,
    color=INK,
    align=PP_ALIGN.LEFT,
    bold=False,
    spacing=None,
    line_gap=1.15,
    anchor=MSO_ANCHOR.TOP,
    margin=0.06,
):
    if isinstance(lines, str):
        lines = [lines]
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = anchor
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(2)
        pPr = p._p.get_or_add_pPr()
        lnSpc = pPr.find(qn("a:lnSpc"))
        if lnSpc is None:
            lnSpc = etree.SubElement(pPr, qn("a:lnSpc"))
        spcPct = lnSpc.find(qn("a:spcPct"))
        if spcPct is None:
            spcPct = etree.SubElement(lnSpc, qn("a:spcPct"))
        spcPct.set("val", str(int(line_gap * 100000)))
        run = p.add_run()
        run.text = line
        set_run_font(run, font, size, color, bold=bold, spacing=spacing)
    return box


def add_rect(slide, l, t, w, h, fill, line=None, line_w=1.0, rounded=False, adj=0.08):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    if rounded:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    sh.shadow.inherit = False
    return sh


def add_picture(slide, path: Path, l, t, w, h):
    return slide.shapes.add_picture(str(path), Inches(l), Inches(t), Inches(w), Inches(h))


def send_to_back(slide, shape):
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def set_bg(slide, image_path: Path):
    pic = slide.shapes.add_picture(str(image_path), Emu(0), Emu(0), SLIDE_W, SLIDE_H)
    send_to_back(slide, pic)
    return pic


def add_left_bar(slide):
    add_rect(slide, 0, 0, 0.09, H_IN, CINNABAR)


def add_header(slide, kicker: str, title: str, assets: dict):
    add_text(slide, 0.55, 0.22, 8.5, 0.32, kicker, font=FONT_BODY, size=12, color=CINNABAR, bold=True, spacing=180)
    add_text(slide, 0.50, 0.46, 11.5, 0.62, title, font=FONT_TITLE, size=32, color=INK, bold=True, spacing=80)
    add_picture(slide, assets["stroke"], 0.55, 1.08, 3.6, 0.09)


def add_footer(slide, page: int, total: int = 12):
    add_text(slide, 0.55, 7.12, 6, 0.28, "纸上沉浮  ·  第一课", font=FONT_BODY, size=11, color=INK_SOFT)
    add_text(
        slide,
        10.6,
        7.12,
        2.2,
        0.28,
        f"{page:02d}  /  {total:02d}",
        font=FONT_BODY,
        size=11,
        color=INK_SOFT,
        align=PP_ALIGN.RIGHT,
    )


def add_card(slide, l, t, w, h, accent=CINNABAR):
    card = add_rect(slide, l, t, w, h, CARD, line=BORDER, line_w=1.15, rounded=True, adj=0.06)
    add_rect(slide, l, t, 0.09, h, accent, rounded=False)
    return card


def add_chip(slide, l, t, w, h, text, fill, font_size=12):
    add_rect(slide, l, t, w, h, fill, rounded=True, adj=0.4)
    add_text(
        slide,
        l,
        t + 0.02,
        w,
        h - 0.02,
        text,
        font=FONT_BODY,
        size=font_size,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.02,
    )


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def slide_cover(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["cover"])
    add_left_bar(s)
    add_picture(s, assets["seal"], 11.55, 0.38, 1.15, 1.15)
    add_text(s, 0.85, 1.55, 10, 0.36, "校本课  ·  文人专题", font=FONT_BODY, size=14, color=CINNABAR, bold=True, spacing=280)
    add_text(s, 0.78, 2.00, 11.5, 1.05, "纸上沉浮", font=FONT_TITLE, size=60, color=INK, bold=True, spacing=420)
    add_text(
        s,
        0.82,
        3.12,
        11,
        0.5,
        "穿越历史，与文人对话",
        font=FONT_TITLE,
        size=26,
        color=INK_SOFT,
        spacing=220,
    )
    add_picture(s, assets["stroke"], 0.85, 3.72, 4.2, 0.10)
    add_text(
        s,
        0.82,
        4.00,
        11,
        0.48,
        "第一课  ·  课程启航，初识诗仙李白",
        font=FONT_BODY,
        size=20,
        color=INK,
        spacing=80,
    )


def slide_intro(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["inner"])
    add_left_bar(s)
    add_header(s, "课程开篇", "课程引言", assets)
    # 两句题记
    add_card(s, 0.55, 1.38, 12.2, 1.55, CINNABAR)
    add_text(s, 0.85, 1.50, 11.6, 0.55, "纸载千年事，墨存万古心。", font=FONT_TITLE, size=24, color=INK, spacing=80)
    add_text(
        s,
        0.85,
        2.08,
        11.6,
        0.7,
        "山河会旧，岁月会走。文人留在纸上的温度、风骨、坦荡与坚守，永远不会消失。",
        font=FONT_BODY,
        size=16,
        color=INK_SOFT,
        line_gap=1.25,
    )
    # 浮 / 沉
    add_card(s, 0.55, 3.15, 5.95, 2.55, CINNABAR)
    add_chip(s, 0.82, 3.38, 0.7, 0.34, "浮", CINNABAR, 14)
    add_text(s, 1.65, 3.36, 4.5, 0.4, "春风得意  ·  少年意气", font=FONT_TITLE, size=18, color=INK)
    add_text(
        s,
        0.82,
        3.90,
        5.4,
        1.55,
        "人生高光、志向勃发、被看见的时刻。\n比如：被赏识、被重用、意气风发。\n古人也会有“走得顺”的时候。",
        font=FONT_BODY,
        size=16,
        color=INK_SOFT,
        line_gap=1.28,
    )
    add_card(s, 6.80, 3.15, 5.95, 2.55, INDIGO)
    add_chip(s, 7.07, 3.38, 0.7, 0.34, "沉", INDIGO, 14)
    add_text(s, 7.90, 3.36, 4.5, 0.4, "失意困顿  ·  逆境坚守", font=FONT_TITLE, size=18, color=INK)
    add_text(
        s,
        7.07,
        3.90,
        5.4,
        1.55,
        "被误解、被贬谪、走投无路的时刻。\n比如：被排挤、被流放，却依旧写诗。\n真正的风骨，往往写在低谷里。",
        font=FONT_BODY,
        size=16,
        color=INK_SOFT,
        line_gap=1.28,
    )
    add_text(
        s,
        0.55,
        5.82,
        12.2,
        1.1,
        "古人的一生有起有落，这些起落都写在了纸上。这门课，就是透过文字看见他们的浮与沉，也映照我们自己的少年成长。",
        font=FONT_BODY,
        size=16,
        color=INK,
        line_gap=1.25,
    )
    add_footer(s, 2)


def slide_goals(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["inner"])
    add_left_bar(s)
    add_header(s, "学习路径", "这门课我们会做什么", assets)
    items = [
        ("01", "读", "读懂人生起落", "看文人的高光，也看他们的低谷。\n抓住一条主线：浮与沉。", CINNABAR),
        ("02", "品", "品文字里的风骨", "不只记诗句，更要读心境。\n温度、坦荡、坚守，都在字里。", INDIGO),
        ("03", "作", "动手做出作品", "文人小卡、翻翻书、A4 海报、短文。\n每节课都有实实在在的产出。", MOSS),
        ("04", "展", "期末集中展示", "班级板报 + 作品集。\n把一学期的沉浮，摊在纸上。", OCHRE),
    ]
    positions = [(0.55, 1.40), (6.80, 1.40), (0.55, 3.85), (6.80, 3.85)]
    for (num, key, title, body, color), (x, y) in zip(items, positions):
        add_card(s, x, y, 5.95, 2.22, color)
        add_chip(s, x + 0.28, y + 0.28, 0.62, 0.34, num, color, 12)
        add_text(s, x + 1.05, y + 0.26, 1.0, 0.4, key, font=FONT_TITLE, size=20, color=color, bold=True)
        add_text(s, x + 2.00, y + 0.28, 3.6, 0.4, title, font=FONT_TITLE, size=20, color=INK)
        add_text(s, x + 0.28, y + 0.85, 5.4, 1.15, body, font=FONT_BODY, size=15, color=INK_SOFT, line_gap=1.25)
    add_text(
        s,
        0.55,
        6.22,
        12.2,
        0.7,
        "这门课不追求背诵大量知识点。认识人、读懂心、做出作品，比记条文更重要。",
        font=FONT_BODY,
        size=15,
        color=INK,
    )
    add_footer(s, 3)


def slide_teacher(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["inner"])
    add_left_bar(s)
    add_header(s, "课程说明", "授课教师：周艳颖", assets)
    add_card(s, 0.55, 1.40, 12.2, 1.35, CINNABAR)
    add_text(
        s,
        0.85,
        1.62,
        11.6,
        0.45,
        "以笔墨润心，以古人养气，以温柔育人",
        font=FONT_TITLE,
        size=24,
        color=INK,
        align=PP_ALIGN.CENTER,
        spacing=120,
    )
    add_text(
        s,
        0.85,
        2.12,
        11.6,
        0.4,
        "借古人的笔墨，在少年时代收获一份沉静、坦荡与坚韧",
        font=FONT_BODY,
        size=16,
        color=INK_SOFT,
        align=PP_ALIGN.CENTER,
    )
    points = [
        ("不是为了多背诗", "诗词是入口，不是终点。我们要看见写诗的那个人。"),
        ("不赶进度", "一节课不塞太多。分课时慢慢走进每一位文人。"),
        ("重在沉淀", "每一节课带走一点思考、一点感悟，比赶完流程更重要。"),
    ]
    for i, (title, body) in enumerate(points):
        x = 0.55 + i * 4.15
        add_card(s, x, 3.00, 3.95, 2.70, CINNABAR if i == 0 else (INDIGO if i == 1 else MOSS))
        add_chip(s, x + 0.28, 3.24, 0.55, 0.32, f"0{i+1}", CINNABAR if i == 0 else (INDIGO if i == 1 else MOSS), 12)
        add_text(s, x + 0.28, 3.72, 3.4, 0.7, title, font=FONT_TITLE, size=20, color=INK)
        add_text(s, x + 0.28, 4.45, 3.4, 1.0, body, font=FONT_BODY, size=15, color=INK_SOFT, line_gap=1.3)
    add_footer(s, 4)


def slide_rules(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["inner"])
    add_left_bar(s)
    add_header(s, "共同约定", "校本课课堂公约", assets)
    rules = [
        ("1", "静心", "入课即安。课堂保持沉静，不随意闲谈打闹。", CINNABAR),
        ("2", "尊重", "尊重笔墨，尊重古人，尊重同学，尊重课堂。", INDIGO),
        ("3", "沉浸", "认真参与小组活动，不敷衍应付。", MOSS),
        ("4", "收获", "每一节课，带走一点思考，一点感悟。", OCHRE),
    ]
    for i, (num, title, body, color) in enumerate(rules):
        y = 1.38 + i * 0.95
        add_card(s, 0.55, y, 12.2, 0.86, color)
        add_chip(s, 0.82, y + 0.24, 0.48, 0.38, num, color, 14)
        add_text(s, 1.50, y + 0.18, 2.2, 0.5, title, font=FONT_TITLE, size=22, color=INK, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, 3.70, y + 0.18, 8.7, 0.52, body, font=FONT_BODY, size=17, color=INK_SOFT, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, 0.55, 5.28, 12.2, 1.55, CARD_WARM, line=BORDER, line_w=1.15, rounded=True, adj=0.06)
    add_text(s, 0.82, 5.42, 11.6, 0.35, "每节课请自备", font=FONT_BODY, size=13, color=CINNABAR, bold=True)
    add_text(
        s,
        0.82,
        5.82,
        11.6,
        0.8,
        "黑色签字笔　　彩笔 / 荧光笔　　校本专用文件夹或手账本",
        font=FONT_TITLE,
        size=20,
        color=INK,
    )
    add_footer(s, 5)


def slide_groups(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["inner"])
    add_left_bar(s)
    add_header(s, "合作方式", "小组分工（四人一组）", assets)
    add_text(
        s,
        0.55,
        1.28,
        12.2,
        0.4,
        "沿用班级四人小组。四个岗位，人人有任务，互相配合。",
        font=FONT_BODY,
        size=16,
        color=INK_SOFT,
    )
    roles = [
        ("记录员", "整理文人资料", "摘抄诗句、梳理要点，把听到的、读到的变成小组共同的笔记。", CINNABAR),
        ("素材员", "搜集人物故事", "找小故事、找金句，为后面的手卡和海报备好素材。", INDIGO),
        ("创作员", "绘画 · 排版 · 手作", "负责小卡、翻翻书、海报的画面与排版，把文字做成作品。", MOSS),
        ("展示员", "分享与汇报", "代表小组发言，把作品和思考讲清楚、讲完整。", OCHRE),
    ]
    for i, (title, sub, body, color) in enumerate(roles):
        x = 0.55 + i * 3.15
        add_card(s, x, 1.78, 3.00, 4.50, color)
        add_chip(s, x + 0.22, 2.00, 0.7, 0.34, f"0{i+1}", color, 12)
        add_text(s, x + 0.18, 2.48, 2.64, 0.62, title, font=FONT_TITLE, size=24, color=INK, align=PP_ALIGN.CENTER)
        add_text(s, x + 0.18, 3.12, 2.64, 0.48, sub, font=FONT_BODY, size=14, color=color, align=PP_ALIGN.CENTER, bold=True)
        add_text(s, x + 0.22, 3.70, 2.56, 2.3, body, font=FONT_BODY, size=14, color=INK_SOFT, align=PP_ALIGN.LEFT, line_gap=1.35)
    add_footer(s, 6)


def slide_libai_meet(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["inner"])
    add_left_bar(s)
    add_header(s, "本课人物", "初识诗仙李白", assets)
    add_text(s, 0.55, 1.28, 12.2, 0.4, "人生有浮亦有沉", font=FONT_TITLE, size=22, color=CINNABAR, spacing=160)
    add_card(s, 0.55, 1.82, 7.55, 4.85, CINNABAR)
    add_text(s, 0.85, 2.05, 7.0, 0.4, "为什么第一个走进李白？", font=FONT_TITLE, size=20, color=INK)
    add_text(
        s,
        0.85,
        2.55,
        6.95,
        3.8,
        "李白的一生，是非常典型的「沉浮」。\n\n少年意气、仗剑远游、供奉翰林，是浮；\n遭排挤、被流放、风骨不改，是沉。\n\n今天只做初步认识——\n故事与诗词，后面用 2—3 节课慢慢精讲，不深挖。",
        font=FONT_BODY,
        size=17,
        color=INK_SOFT,
        line_gap=1.28,
    )
    facts = [("字", "太白"), ("号", "青莲居士"), ("人称", "诗仙"), ("时代", "盛唐")]
    for i, (k, v) in enumerate(facts):
        y = 1.82 + i * 1.22
        add_card(s, 8.35, y, 4.40, 1.10, INDIGO)
        add_text(s, 8.55, y + 0.12, 4.0, 0.32, k, font=FONT_BODY, size=12, color=CINNABAR, bold=True)
        add_text(s, 8.55, y + 0.42, 4.0, 0.5, v, font=FONT_TITLE, size=22, color=INK)
    add_footer(s, 7)


def slide_life(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["inner"])
    add_left_bar(s)
    add_header(s, "人物轮廓", "李白的人生轮廓（浮与沉）", assets)
    stages = [
        ("01", "少年浮", "浮", "天资卓绝\n少年意气\n胸怀壮志", CINNABAR),
        ("02", "青年游", "游", "仗剑远游\n遍览山河\n肆意洒脱", MOSS),
        ("03", "中年浮", "浮", "入仕长安\n供奉翰林\n人生高光", OCHRE),
        ("04", "晚年沉", "沉", "遭排挤流放\n历经磨难\n风骨不改", INDIGO),
    ]
    for i, (num, title, tag, body, color) in enumerate(stages):
        x = 0.55 + i * 3.15
        add_card(s, x, 1.48, 3.00, 3.70, color)
        add_chip(s, x + 0.22, 1.70, 0.62, 0.32, num, color, 12)
        add_chip(s, x + 2.05, 1.70, 0.70, 0.32, tag, color, 12)
        add_text(s, x + 0.12, 2.16, 2.76, 0.58, title, font=FONT_TITLE, size=24, color=INK, align=PP_ALIGN.CENTER)
        add_text(s, x + 0.38, 2.85, 2.40, 2.1, body, font=FONT_BODY, size=18, color=INK_SOFT, align=PP_ALIGN.LEFT, line_gap=1.55)
    add_text(
        s,
        0.55,
        5.38,
        12.2,
        0.5,
        "先记住这条人生线。故事与诗，后面几课再展开。",
        font=FONT_BODY,
        size=16,
        color=INK,
    )
    add_footer(s, 8)


def slide_poems(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["inner"])
    add_left_bar(s)
    add_header(s, "读一句，见底色", "李白的人生态度", assets)
    poems = [
        ("长风破浪会有时，", "直挂云帆济沧海。", "逆境不屈", "逆境里，不低头。\n再难的水路，也要扬帆过去。", CINNABAR),
        ("天生我材必有用，", "千金散尽还复来。", "少年自信", "才干在，路就在。\n豪气在，就不怕从头来。", INDIGO),
    ]
    for i, (l1, l2, tag, note, color) in enumerate(poems):
        y = 1.40 + i * 2.35
        add_card(s, 0.55, y, 12.2, 2.20, color)
        add_chip(s, 0.85, y + 0.28, 1.35, 0.36, tag, color, 13)
        add_text(s, 0.85, y + 0.78, 6.6, 0.48, l1, font=FONT_POETRY, size=26, color=INK, spacing=80)
        add_text(s, 0.85, y + 1.30, 6.6, 0.48, l2, font=FONT_POETRY, size=26, color=INK, spacing=80)
        add_rect(s, 7.70, y + 0.55, 0.015, 1.30, PAPER_LINE)
        add_text(s, 8.00, y + 0.70, 4.4, 1.15, note, font=FONT_BODY, size=16, color=INK_SOFT, line_gap=1.35)
    add_text(
        s,
        0.55,
        6.15,
        12.2,
        0.7,
        "一句写逆境里不低头，一句写少年的自信。这就是李白的底色。后面几课，我们会细细品读这两句，以及更多诗作。",
        font=FONT_BODY,
        size=15,
        color=INK_SOFT,
        line_gap=1.2,
    )
    add_footer(s, 9)


def slide_homework(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["inner"])
    add_left_bar(s)
    add_header(s, "本课延伸", "课后小组任务：搜集李白素材", assets)
    tasks = [
        ("任务一", "诗句", "搜集李白 3—5 句经典诗句，抄写整理。", "记录员主责，全组共读。"),
        ("任务二", "故事", "简单搜集 1—2 个李白小故事，写短即可。", "素材员主责，创作员可配小插画。"),
        ("任务三", "印象", "你印象中的李白是什么样子？写 2—3 句话。", "每人先写，展示员课上汇总。"),
    ]
    for i, (label, key, body, hint) in enumerate(tasks):
        y = 1.40 + i * 1.70
        color = [CINNABAR, INDIGO, MOSS][i]
        add_card(s, 0.55, y, 12.2, 1.55, color)
        add_chip(s, 0.85, y + 0.28, 1.15, 0.36, label, color, 12)
        add_text(s, 2.20, y + 0.26, 2.2, 0.42, key, font=FONT_TITLE, size=22, color=INK)
        add_text(s, 0.85, y + 0.78, 11.6, 0.55, body, font=FONT_BODY, size=17, color=INK)
        add_text(s, 8.3, y + 0.28, 4.1, 0.4, hint, font=FONT_BODY, size=13, color=INK_SOFT, align=PP_ALIGN.RIGHT)
    add_footer(s, 10)


def slide_preview(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["inner"])
    add_left_bar(s)
    add_header(s, "往前看", "后续课程预告", assets)
    items = [
        ("第 2—3 课", "深度走进李白", "故事、诗词、人物精神。把今天的轮廓，慢慢画清楚。", CINNABAR),
        ("之后", "走进杜甫", "感受另一种文人的沉浮。诗仙之后，是诗圣。", INDIGO),
        ("学期产出", "作品留在纸上", "文人小卡、翻翻书、A4 海报、感悟短文、班级板报素材。", MOSS),
    ]
    for i, (when, title, body, color) in enumerate(items):
        y = 1.42 + i * 1.72
        add_card(s, 0.55, y, 12.2, 1.55, color)
        add_chip(s, 0.85, y + 0.55, 2.15, 0.42, when, color, 13)
        add_text(s, 3.25, y + 0.28, 9.1, 0.5, title, font=FONT_TITLE, size=22, color=INK)
        add_text(s, 3.25, y + 0.82, 9.1, 0.5, body, font=FONT_BODY, size=16, color=INK_SOFT)
    add_footer(s, 11)


def slide_close(prs, assets):
    s = blank_slide(prs)
    set_bg(s, assets["close"])
    add_left_bar(s)
    add_text(s, 0.85, 1.55, 11.5, 0.4, "课堂小结", font=FONT_BODY, size=14, color=CINNABAR, bold=True, spacing=280)
    add_text(
        s,
        0.78,
        2.15,
        11.8,
        1.1,
        "一纸山河，万古温柔",
        font=FONT_TITLE,
        size=40,
        color=INK,
        spacing=180,
    )
    add_picture(s, assets["stroke"], 0.85, 3.35, 4.0, 0.10)
    add_text(
        s,
        0.82,
        3.70,
        11.5,
        1.4,
        "愿我们在笔墨沉浮之中，\n见天地，见众生，见自己。",
        font=FONT_TITLE,
        size=24,
        color=INK_SOFT,
        line_gap=1.45,
        spacing=60,
    )
    add_text(s, 0.82, 5.55, 11.5, 0.5, "下节课，我们真正走进李白。", font=FONT_BODY, size=16, color=INK)
    add_picture(s, assets["seal"], 11.55, 5.85, 1.05, 1.05)
    add_footer(s, 12)


def build() -> Path:
    assets = build_assets()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide_cover(prs, assets)
    slide_intro(prs, assets)
    slide_goals(prs, assets)
    slide_teacher(prs, assets)
    slide_rules(prs, assets)
    slide_groups(prs, assets)
    slide_libai_meet(prs, assets)
    slide_life(prs, assets)
    slide_poems(prs, assets)
    slide_homework(prs, assets)
    slide_preview(prs, assets)
    slide_close(prs, assets)
    prs.save(str(OUT_PPTX))
    return OUT_PPTX


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
